# -*- coding: utf-8 -*-
"""Word (.docx) 生成引擎：封面、标题、段落、列表、表格、批注框、页脚等可复用组件。"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from design import FONT_CN_HEAD, FONT_CN_BODY, FONT_EN, theme, rgb


def _set_cell_bg(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcpr.append(shd)


def _set_para_shading(paragraph, hex_color):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    ppr.append(shd)


def _set_bottom_border(paragraph, color, size=12, space=4):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), str(space))
    bottom.set(qn('w:color'), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def _set_left_border(paragraph, color, size=24, space=6):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(size))
    left.set(qn('w:space'), str(space))
    left.set(qn('w:color'), color)
    pbdr.append(left)
    ppr.append(pbdr)


def _set_table_borders(table, color="BFBFBF", size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(size))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        borders.append(e)
    tblPr.append(borders)


def _run(paragraph, text, font=FONT_CN_BODY, size=11, color="222222",
         bold=False, italic=False, en_font=FONT_EN):
    r = paragraph.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))
    rpr = r._r.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font)
    rfonts.set(qn('w:ascii'), en_font)
    rfonts.set(qn('w:hAnsi'), en_font)
    return r


def _page_field(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def new_document(th):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    sec.left_margin = Inches(0.95)
    sec.right_margin = Inches(0.95)
    st = doc.styles['Normal']
    st.font.name = FONT_CN_BODY
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), FONT_CN_BODY)
    rfonts.set(qn('w:ascii'), FONT_EN)
    rfonts.set(qn('w:hAnsi'), FONT_EN)
    return doc, sec


def add_cover(doc, th, title, subtitle, meta=None, brand="寄海文库 · 模板中心"):
    # 顶部主色横幅
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    _set_cell_bg(cell, th["primary"])
    cell.width = Inches(6.6)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(26)
    _run(p, title, font=FONT_CN_HEAD, size=30, color="FFFFFF", bold=True)
    if subtitle:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(6)
        _run(p2, subtitle, font=FONT_EN, size=12, color="EAF1F8", bold=False)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    p2.paragraph_format.space_after = Pt(26)
    # 控制横幅高度
    tr = cell._tc.get_or_add_tcPr()
    h = OxmlElement('w:tcH'); h.set(qn('w:val'), 'exact'); h.set(qn('w:h'), '2400')
    tr.append(h)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 金色细线
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_bottom_border(line, th["gold"], size=18, space=0)
    sp = line.paragraph_format; sp.space_before = Pt(2); sp.space_after = Pt(14)

    if meta:
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, m in enumerate(meta):
            if i:
                meta_p.add_run("      ")
            _run(meta_p, m, font=FONT_CN_BODY, size=10.5, color=th["muted"])
        meta_p.paragraph_format.space_after = Pt(18)


def add_h1(doc, th, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    _set_bottom_border(p, th["accent"], size=8, space=4)
    _run(p, "▍ ", font=FONT_CN_HEAD, size=15, color=th["accent"], bold=True)
    _run(p, text, font=FONT_CN_HEAD, size=15, color=th["primary"], bold=True)
    return p


def add_h2(doc, th, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    _run(p, text, font=FONT_CN_HEAD, size=12.5, color=th["dark"], bold=True)
    return p


def add_para(doc, th, text, size=11, color=None, bold=False, align="left",
             space_after=6, italic=False):
    p = doc.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.4
    _run(p, text, font=FONT_CN_BODY, size=size, color=color or th["text"],
         bold=bold, italic=italic)
    return p


def add_bullets(doc, th, items, size=11):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.left_indent = Pt(18)
        _run(p, "● ", font=FONT_CN_BODY, size=size, color=th["accent"], bold=True)
        _run(p, it, font=FONT_CN_BODY, size=size, color=th["text"])


def add_callout(doc, th, text, label=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    _set_cell_bg(cell, th["light"])
    _set_left_border(cell.paragraphs[0], th["accent"], size=36, space=8)
    cell.width = Inches(6.5)
    cell.paragraphs[0].paragraph_format.space_before = Pt(6)
    cell.paragraphs[0].paragraph_format.space_after = Pt(6)
    if label:
        _run(cell.paragraphs[0], label + "  ", font=FONT_CN_HEAD, size=10.5,
             color=th["primary"], bold=True)
    _run(cell.paragraphs[0], text, font=FONT_CN_BODY, size=10.5, color=th["text"])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, th, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        _set_cell_bg(hdr[i], th["primary"])
        if widths:
            hdr[i].width = Inches(widths[i])
        hp = hdr[i].paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(hp, h, font=FONT_CN_HEAD, size=10.5, color="FFFFFF", bold=True)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row):
            if widths:
                cells[ci].width = Inches(widths[ci])
            cp = cells[ci].paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci > 0 else WD_ALIGN_PARAGRAPH.LEFT
            if ri % 2 == 1:
                _set_cell_bg(cells[ci], "F2F5F9")
            _run(cp, str(val), font=FONT_CN_BODY, size=10, color=th["text"])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_footer(doc, th, brand):
    sec = doc.sections[0]
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, brand + "    ", font=FONT_EN, size=9, color=th["muted"])
    _run(p, "第 ", font=FONT_EN, size=9, color=th["muted"])
    _page_field(p)
    _run(p, " 页", font=FONT_EN, size=9, color=th["muted"])


def render_word(spec, out_path):
    th = theme(spec.get("theme", 0))
    doc, sec = new_document(th)
    add_cover(doc, th, spec["title"], spec.get("subtitle", ""),
              spec.get("cover_meta"), spec.get("brand", "寄海文库 · 模板中心"))
    for blk in spec.get("blocks", []):
        t = blk["type"]
        if t == "h1":
            add_h1(doc, th, blk["text"])
        elif t == "h2":
            add_h2(doc, th, blk["text"])
        elif t == "para":
            add_para(doc, th, blk["text"], size=blk.get("size", 11),
                     align=blk.get("align", "left"), bold=blk.get("bold", False),
                     color=blk.get("color"), space_after=blk.get("space_after", 6))
        elif t == "bullets":
            add_bullets(doc, th, blk["items"], size=blk.get("size", 11))
        elif t == "callout":
            add_callout(doc, th, blk["text"], blk.get("label"))
        elif t == "table":
            add_table(doc, th, blk["headers"], blk["rows"], blk.get("widths"))
        elif t == "spacer":
            doc.add_paragraph().paragraph_format.space_after = Pt(blk.get("h", 10))
        elif t == "note":
            add_para(doc, th, blk["text"], size=9.5, color=th["muted"], italic=True,
                     align="left", space_after=4)
    add_footer(doc, th, spec.get("brand", "寄海文库 · 模板中心"))
    doc.save(out_path)
